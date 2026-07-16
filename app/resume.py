"""Résumé parsing + skill matching (Phase 5).

parse_resume(filename, data) -> plain text from PDF / DOCX / TXT.
heuristic_profile(text) -> {skills, titles, seniority, summary} without an API key.
matched_skills(text, skills) -> which of `skills` appear (word-boundary) in `text`.
"""
from __future__ import annotations

import io
import re

MAX_RESUME_BYTES = 5 * 1024 * 1024      # 5 MB upload cap
MAX_RESUME_CHARS = 20000                # store/scan at most this much text
MAX_DECOMPRESSED_BYTES = 60 * 1024 * 1024  # reject decompression bombs (docx = zip)
MAX_PDF_PAGES = 100                     # bound PDF extraction work


class ResumeError(RuntimeError):
    """Raised when a résumé file can't be read or has no extractable text."""


def _capped_join(chunks) -> str:
    """Join text chunks, stopping once MAX_RESUME_CHARS is reached (bounds memory)."""
    out: list[str] = []
    total = 0
    for chunk in chunks:
        chunk = chunk or ""
        out.append(chunk)
        total += len(chunk)
        if total >= MAX_RESUME_CHARS:
            break
    return "\n".join(out)


def parse_resume(filename: str, data: bytes) -> str:
    """Extract plain text from an uploaded résumé (PDF / DOCX / TXT / MD)."""
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))

            def _pdf_pages():
                for i, page in enumerate(reader.pages):
                    if i >= MAX_PDF_PAGES:
                        break
                    yield page.extract_text() or ""

            text = _capped_join(_pdf_pages())
        elif name.endswith(".docx"):
            import zipfile

            # A .docx is a ZIP — reject decompression bombs before parsing the XML.
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if sum(zi.file_size for zi in zf.infolist()) > MAX_DECOMPRESSED_BYTES:
                    raise ResumeError("Document is too large when decompressed.")

            import docx

            doc = docx.Document(io.BytesIO(data))

            def _docx_parts():
                for p in doc.paragraphs:
                    yield p.text
                for table in doc.tables:  # skills are often in tables
                    for row in table.rows:
                        for cell in row.cells:
                            yield cell.text

            text = _capped_join(_docx_parts())
        elif name.endswith((".txt", ".md")):
            text = data.decode("utf-8", errors="ignore")
        elif name.endswith(".doc"):
            raise ResumeError("Legacy .doc isn't supported — save as PDF or .docx.")
        else:
            # Best-effort: treat unknown types as text.
            text = data.decode("utf-8", errors="ignore")
    except ResumeError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ResumeError(f"Couldn't read the file: {exc}") from exc

    text = (text or "").strip()
    if not text:
        raise ResumeError("No text could be extracted (is it a scanned image PDF?).")
    return text[:MAX_RESUME_CHARS]


# A pragmatic vocabulary of tech skills for the no-API-key fallback. Multi-word
# entries are matched as phrases. Kept lowercase.
SKILLS_VOCAB = [
    "python", "javascript", "typescript", "php", "java", "c#", "c++", "go", "golang",
    "ruby", "rust", "kotlin", "swift", "scala", "perl", "bash", "sql", "html", "css",
    "laravel", "django", "flask", "fastapi", "express", "node.js", "node", "nestjs",
    "react", "react native", "next.js", "nextjs", "vue", "nuxt", "angular", "svelte",
    "tailwind", "bootstrap", "jquery", "redux",
    "rest", "rest api", "graphql", "grpc", "websockets", "microservices",
    "mysql", "postgresql", "postgres", "mongodb", "redis", "sqlite", "elasticsearch",
    "mariadb", "dynamodb", "oracle", "mssql",
    "aws", "gcp", "azure", "docker", "kubernetes", "terraform", "ansible", "nginx",
    "linux", "git", "ci/cd", "jenkins", "github actions",
    "ai", "machine learning", "deep learning", "nlp", "llm", "openai", "claude",
    "langchain", "tensorflow", "pytorch", "pandas", "numpy", "scikit-learn",
    "data analysis", "data engineering", "etl",
    "chatbot", "integration", "dashboard", "saas", "api", "asterisk", "voip", "webrtc",
    "sentiment analysis", "speech-to-text", "transcription", "real-time",
    "wordpress", "shopify", "magento", "woocommerce",
    "agile", "scrum", "rbac", "oauth", "stripe", "twilio", "celery", "rabbitmq", "kafka",
]

_SENIOR_RE = re.compile(r"\b(lead|principal|staff|head of|architect|director)\b", re.I)
_MID_RE = re.compile(r"\b(senior|sr\.?)\b", re.I)
_JUNIOR_RE = re.compile(r"\b(junior|jr\.?|intern|entry[- ]level|graduate)\b", re.I)


def _guess_seniority(text: str) -> str:
    if _SENIOR_RE.search(text):
        return "lead"
    if _MID_RE.search(text):
        return "senior"
    if _JUNIOR_RE.search(text):
        return "junior"
    return "mid"


def _skill_pattern(term: str) -> str:
    """Boundary-aware pattern. Only require a word boundary on a side that ends in
    an alphanumeric — so skills ending in punctuation (c++, c#) still match 'C++11'.
    """
    escaped = re.escape(term)
    start = r"(?<!\w)" if term[:1].isalnum() else ""
    end = r"(?!\w)" if term[-1:].isalnum() else ""
    return start + escaped + end


def extract_skills_heuristic(text: str) -> list[str]:
    """Match the skills vocabulary against résumé text (word-boundary, deduped)."""
    low = (text or "").lower()
    found: list[str] = []
    seen: set[str] = set()
    for skill in SKILLS_VOCAB:
        if skill in seen:
            continue
        if re.search(_skill_pattern(skill), low):
            seen.add(skill)
            found.append(skill)
    return found


def heuristic_profile(text: str) -> dict:
    """Build a profile without an API key (skills vocab + a text-snippet summary)."""
    skills = extract_skills_heuristic(text)
    summary = " ".join(text.split())[:400]
    return {
        "skills": skills,
        "titles": [],
        "seniority": _guess_seniority(text),
        "summary": summary,
    }


def matched_skills(text: str, skills: list[str]) -> list[str]:
    """Which of `skills` appear as whole words/phrases in `text` (deduped, ordered)."""
    low = (text or "").lower()
    out: list[str] = []
    seen: set[str] = set()
    for raw in skills or []:
        skill = (raw or "").strip().lower()
        if not skill or skill in seen:
            continue
        if re.search(_skill_pattern(skill), low):
            seen.add(skill)
            out.append(skill)
    return out
